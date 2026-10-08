"""Genera los archivos descargables de una fuente: ZIP de CSV, libro Excel y script PostgreSQL."""

from __future__ import annotations

import io
import zipfile
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TIPO_SQL = {"texto": "TEXT", "entero": "BIGINT", "booleano": "BOOLEAN", "fecha": "DATE"}
LOTE_INSERT = 500


def _valor(valor, tipo: str):
    """Convierte un valor del DataFrame al tipo Python correspondiente a la columna del modelo."""
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        return None
    if tipo == "entero":
        return int(valor)
    if tipo == "booleano":
        return bool(valor) if not isinstance(valor, str) else valor == "True"
    if tipo == "fecha":
        return date.fromisoformat(str(valor)[:10])
    return str(valor)


def _filas(df: pd.DataFrame, columnas: list[dict]):
    tipos = [c["tipo"] for c in columnas]
    for registro in df.itertuples(index=False, name=None):
        yield [_valor(v, t) for v, t in zip(registro, tipos)]


def _orden_por_dependencias(modelo: dict) -> list[str]:
    orden, visitadas = [], set()

    def visitar(tabla: str) -> None:
        if tabla in visitadas:
            return
        visitadas.add(tabla)
        for c in modelo[tabla]["columnas"]:
            if c["fk"]:
                visitar(c["fk"].split(".")[0])
        orden.append(tabla)

    for tabla in modelo:
        visitar(tabla)
    return orden


# ---------------------------------------------------------------- CSV (ZIP)

def exportar_zip(destino: Path, carpeta_csv: Path, tablas: dict[str, pd.DataFrame], modelo: dict) -> None:
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre in tablas:
            subcarpeta = "origen/" if modelo[nombre].get("origen") else ""
            z.write(carpeta_csv / f"{nombre}.csv", f"{subcarpeta}{nombre}.csv")
        z.write(carpeta_csv / "modelo.json", "modelo.json")


# ---------------------------------------------------------------- Excel

def exportar_excel(destino: Path, tablas: dict[str, pd.DataFrame], modelo: dict, calidad: list[dict],
                   titulo: str) -> None:
    wb = Workbook()
    negrita = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="2F6FED")

    def encabezado(ws, columnas: list[str]) -> None:
        ws.append(columnas)
        for celda in ws[ws.max_row]:
            celda.font, celda.fill = negrita, relleno
            celda.alignment = Alignment(vertical="center")

    def ajustar(ws) -> None:
        for i, columna in enumerate(ws.iter_cols(values_only=True), start=1):
            largo = max((len(str(v)) for v in columna if v is not None), default=8)
            ws.column_dimensions[get_column_letter(i)].width = min(max(largo + 2, 10), 60)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    indice = wb.active
    indice.title = "Indice"
    indice.append([titulo])
    indice["A1"].font = Font(bold=True, size=14)
    indice.append([])
    encabezado(indice, ["Hoja", "Dominio", "Filas", "Descripción"])
    for nombre, meta in modelo.items():
        indice.append([nombre, meta["dominio"], len(tablas[nombre]), meta["descripcion"]])
        indice.cell(indice.max_row, 1).hyperlink = f"#'{nombre}'!A1"
        indice.cell(indice.max_row, 1).style = "Hyperlink"
    for i, ancho in enumerate([26, 22, 10, 70], start=1):
        indice.column_dimensions[get_column_letter(i)].width = ancho

    hoja_calidad = wb.create_sheet("Calidad")
    encabezado(hoja_calidad, ["Tabla", "Columna", "Problema", "Filas afectadas", "Acción"])
    for q in calidad:
        hoja_calidad.append([q["tabla"], q["columna"], q["problema"], q["filas_afectadas"], q["accion"]])
    ajustar(hoja_calidad)

    for nombre, meta in modelo.items():
        ws = wb.create_sheet(nombre[:31])
        encabezado(ws, [c["nombre"] for c in meta["columnas"]])
        for fila in _filas(tablas[nombre], meta["columnas"]):
            ws.append(fila)
        for i, c in enumerate(meta["columnas"], start=1):
            formato = {"fecha": "yyyy-mm-dd", "entero": "#,##0"}.get(c["tipo"])
            if formato:
                for (celda,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    celda.number_format = formato
        ajustar(ws)
    wb.save(destino)


# ---------------------------------------------------------------- PostgreSQL

def _literal(valor) -> str:
    if valor is None:
        return "NULL"
    if isinstance(valor, bool):
        return "TRUE" if valor else "FALSE"
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, date):
        return f"'{valor.isoformat()}'"
    return "'" + str(valor).replace("'", "''") + "'"


def exportar_sql(destino: Path, tablas: dict[str, pd.DataFrame], modelo: dict, esquema: str, titulo: str) -> None:
    destino.write_text(sql_postgres(tablas, modelo, esquema, titulo, destino.name), encoding="utf-8")


def sql_postgres(tablas: dict[str, pd.DataFrame], modelo: dict, esquema: str, titulo: str,
                 archivo: str = "script.sql") -> str:
    out = io.StringIO()
    out.write(f"-- {titulo}\n-- Script para PostgreSQL: esquema, tablas, claves y datos.\n")
    out.write(f"-- Uso: psql -d <base_de_datos> -f {archivo}\n\n")
    out.write("SET client_encoding = 'UTF8';\n\nBEGIN;\n\n")
    out.write(f"DROP SCHEMA IF EXISTS {esquema} CASCADE;\nCREATE SCHEMA {esquema};\nSET search_path TO {esquema};\n\n")

    orden = _orden_por_dependencias(modelo)
    for nombre in orden:
        meta = modelo[nombre]
        definiciones = []
        for c in meta["columnas"]:
            linea = f"    {c['nombre']} {TIPO_SQL.get(c['tipo'], 'TEXT')}"
            if c["pk"]:
                linea += " PRIMARY KEY"
            if c["fk"]:
                tabla_ref, col_ref = c["fk"].split(".")
                linea += f" REFERENCES {tabla_ref} ({col_ref})"
            definiciones.append(linea)
        out.write(f"CREATE TABLE {nombre} (\n" + ",\n".join(definiciones) + "\n);\n")
        out.write(f"COMMENT ON TABLE {nombre} IS {_literal(meta['descripcion'])};\n")
        for c in meta["columnas"]:
            if c["descripcion"]:
                out.write(f"COMMENT ON COLUMN {nombre}.{c['nombre']} IS {_literal(c['descripcion'])};\n")
        out.write("\n")

    for nombre in orden:
        meta = modelo[nombre]
        filas = list(_filas(tablas[nombre], meta["columnas"]))
        if not filas:
            continue
        columnas = ", ".join(c["nombre"] for c in meta["columnas"])
        out.write(f"-- {nombre}: {len(filas)} filas\n")
        for i in range(0, len(filas), LOTE_INSERT):
            valores = ",\n".join("    (" + ", ".join(_literal(v) for v in f) + ")" for f in filas[i:i + LOTE_INSERT])
            out.write(f"INSERT INTO {nombre} ({columnas}) VALUES\n{valores};\n")
        out.write("\n")

    out.write("COMMIT;\n")
    return out.getvalue()


def exportar_todo(carpeta: Path, id_fuente: str, nombre_fuente: str, tablas: dict[str, pd.DataFrame],
                  modelo: dict, calidad: list[dict]) -> dict[str, str]:
    """Escribe los tres formatos en <carpeta>/export y devuelve sus rutas relativas a la carpeta de la fuente."""
    destino = carpeta / "export"
    destino.mkdir(exist_ok=True)
    titulo = f"Lavadero de datos · {nombre_fuente}"
    base = f"lavadero_{id_fuente}"
    archivos = {"csv": f"{base}_csv.zip", "excel": f"{base}.xlsx", "sql": f"{base}_postgres.sql"}
    exportar_zip(destino / archivos["csv"], carpeta, tablas, modelo)
    exportar_excel(destino / archivos["excel"], tablas, modelo, calidad, titulo)
    exportar_sql(destino / archivos["sql"], tablas, modelo, base, titulo)
    return {formato: f"export/{archivo}" for formato, archivo in archivos.items()}
